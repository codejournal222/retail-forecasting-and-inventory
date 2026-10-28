import json
from datetime import datetime
from pathlib import Path

REQUIRED_SHEETS = {
    "Sales",
    "Products",
    "Forecasts",
    "Inventory",
    "Model_Comparison",
    "Summary",
}
SHEET_ALIASES = {"Replenishment": "Inventory", "Future_Forecast": "Forecasts"}


# Export result tables and editable replenishment formulas
def export_excel(output):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.properties import CalcProperties

    out = Path(output)
    data = json.loads((out / "workbook_data.json").read_text())
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    wb.calculation = CalcProperties(
        calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True
    )
    manifest = json.loads((out / "manifest.json").read_text())

    # Keep the most immediately useful result sheets first
    order = [
        "Replenishment",
        "Model_Comparison",
        "Sales",
        "Protection_Comparison",
        "Inventory_Comparison",
        "Future_Forecast",
        "Products",
        "Fold_Metrics",
        "Protection_Forecasts",
        "Inventory_Folds",
        "Validation",
        "Calibration",
        "Cleaning",
        "Cleaning_Sensitivity",
        "Cost_Sensitivity",
        "Test_Windows",
        "Calendar_Review",
        "Calendar_Sensitivity",
        "HW_Reference",
    ]
    for name in order:
        records = data[name]
        ws = wb.create_sheet(SHEET_ALIASES.get(name, name))
        ws.sheet_view.showGridLines = False
        if not records:
            continue
        headers = list(records[0])
        ws.append(headers)
        for record in records:
            # Prevent source descriptions from being interpreted as Excel formulas
            def value(k):
                v = record[k]
                if k in ["Date", "Origin", "Start", "End"] and isinstance(v, str):
                    return datetime.fromisoformat(v.replace("Z", "+00:00")).replace(
                        tzinfo=None
                    )
                return (
                    ("'" + v)
                    if isinstance(v, str) and v.startswith(("=", "+", "-", "@"))
                    else v
                )

            ws.append([value(k) for k in headers])

        # Apply consistent headers, column widths and number formats
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="344B75")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 34
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            key = col[0].value
            width = min(
                48,
                max(
                    16,
                    len(str(key)) + 2,
                    max(len(str(c.value or "")) for c in col[:30]) + 1,
                ),
            )
            ws.column_dimensions[col[0].column_letter].width = width
            for cell in col[1:]:
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd"
                if key in ["FillRate", "SelectionZeroShare"]:
                    cell.number_format = "0.0%"
                if key.endswith("_pct"):
                    cell.number_format = '0.0"%"'
                if "Cost" in key or key == "RevenueGBP":
                    cell.number_format = '"£"#,##0.00'

        # Add formulas and highlight the editable stock inputs
        if name == "Replenishment":
            idx = {x: get_column_letter(i + 1) for i, x in enumerate(headers)}
            for r in range(2, ws.max_row + 1):
                ref = lambda key: idx[key] + str(r)
                ws[ref("ReorderPoint")] = (
                    f"={ref('ForecastLeadUnits')}+{ref('SafetyStock')}"
                )
                ws[ref("OrderUpTo")] = (
                    f"={ref('ForecastProtectionUnits')}+{ref('SafetyStock')}"
                )
                ws[ref("SuggestedOrder")] = (
                    f"=MAX(0,ROUNDUP({ref('OrderUpTo')}-{ref('CurrentStock')}-{ref('OnOrder')},0))"
                )
            for key in ["CurrentStock", "OnOrder"]:
                validation = DataValidation(
                    type="decimal", operator="greaterThanOrEqual", formula1=0
                )
                validation.errorTitle = "Nonnegative stock required"
                validation.error = "Enter zero or a positive number."
                validation.showErrorMessage = True
                validation.errorStyle = "stop"
                validation.allow_blank = False
                ws.add_data_validation(validation)
                validation.add(f"{idx[key]}2:{idx[key]}{ws.max_row}")
                for row in ws.iter_rows(
                    min_row=2,
                    min_col=headers.index(key) + 1,
                    max_col=headers.index(key) + 1,
                ):
                    row[0].fill = PatternFill("solid", fgColor="FFF2D9")
    summary.append(["Retail forecasting and inventory study"])
    summary.append(
        [
            f"{manifest['analysis_start']} to {manifest['analysis_end']} | {len(data['Products'])} selected products"
        ]
    )
    summary.append(
        [
            "Illustrative inventory costs, not realised savings. Edit amber stock cells on Inventory."
        ]
    )
    summary.append(
        ["Model", "Daily MAE", "Protection MAE", "Simulated cost (£)", "Unit fill rate"]
    )
    for r, record in enumerate(data["Model_Comparison"], 5):
        model = record["Model"]
        summary.cell(r, 1, model)
        for col, source, metric in [
            (2, "Model_Comparison", "MAE"),
            (3, "Protection_Comparison", "MAE"),
            (4, "Inventory_Comparison", "TotalCost"),
            (5, "Inventory_Comparison", "FillRate"),
        ]:
            source_row = next(
                i + 2 for i, item in enumerate(data[source]) if item["Model"] == model
            )
            source_col = get_column_letter(list(data[source][0]).index(metric) + 1)
            summary.cell(r, col, f"='{source}'!{source_col}{source_row}")
    summary.append([])
    for text in [
        f"Unresolved non-closed dates without source records: {manifest['calendar_unknown_zero_days']}.",
        "Sales contains selected-product daily units and retained revenue, with source-coverage flags.",
        "Forecasts are after this historical sample, not forecasts for today's date.",
        "Calendar_Sensitivity compares training treatments on the same known test dates only.",
        "Safety stock has only four default calibration origins; no service guarantee is claimed.",
        "Live formulas recalculate in Excel. CSVs provide independently computed initial values.",
    ]:
        summary.append([text])
    for row in [1, 2, 3, *range(10, 16)]:
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    for row in summary:
        for cell in row:
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    summary["A1"].font = Font(name="Calibri", size=21, bold=True, color="344B75")
    for cell in summary[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="344B75")
    for r in range(5, 9):
        summary.cell(r, 2).number_format = "0.00"
        summary.cell(r, 3).number_format = "0.00"
        summary.cell(r, 4).number_format = '"£"#,##0.00'
        summary.cell(r, 5).number_format = "0.0%"
    for col, width in {"A": 29, "B": 22, "C": 22, "D": 25, "E": 22}.items():
        summary.column_dimensions[col].width = width
    for r in range(1, 16):
        summary.row_dimensions[r].height = 32
    summary.row_dimensions[3].height = 42
    summary.freeze_panes = "B5"
    summary.sheet_view.showGridLines = False
    path = out / "Retail_Results.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("folder", nargs="?", default="output")
    print(export_excel(parser.parse_args().folder))
