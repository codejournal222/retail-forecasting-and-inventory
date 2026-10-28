import json
import math
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from calendar_review import past_weekday_fill
from forecasting import GRIDS, hw
from local_excel import export_excel
from reference_checks import reference_hw
from workbook_checks import check_workbook, formula_value
from verify_outputs import verify


def workbook_fixture(folder):
    #Small independent inputs; tests do not depend on distributed output files.
    models = ["Moving average", "Seasonal naive", "Holt-Winters", "Random Forest"]
    data = {
        name: []
        for name in [
            "Fold_Metrics",
            "Protection_Forecasts",
            "Inventory_Folds",
            "Validation",
            "Calibration",
            "Cleaning",
            "Cleaning_Sensitivity",
            "Cost_Sensitivity",
            "Test_Windows",
            "Calendar_Review",
            "Calendar_Sensitivity",
            "HW_Reference",
        ]
    }
    data.update(
        {
            "Products": [{"StockCode": "A", "Description": "=untrusted source text"}],
            "Model_Comparison": [
                {"Model": m, "MAE": i + 1.5} for i, m in enumerate(models)
            ],
            "Protection_Comparison": [
                {"Model": m, "MAE": i + 10.5} for i, m in enumerate(models)
            ],
            "Inventory_Comparison": [
                {"Model": m, "TotalCost": 100 + i, "FillRate": 0.8 + i * 0.01}
                for i, m in enumerate(models)
            ],
            "Future_Forecast": [
                {"StockCode": "A", "Date": "2012-01-01", "Forecast": 10.5}
            ],
            "Sales": [
                {
                    "StockCode": "A",
                    "Date": "2011-12-31",
                    "Units": 10,
                    "RevenueGBP": 20.0,
                }
            ],
            "Replenishment": [
                {
                    "StockCode": "A",
                    "ForecastLeadUnits": 50.5,
                    "ForecastProtectionUnits": 100.5,
                    "SafetyStock": 10.0,
                    "CurrentStock": 20.0,
                    "OnOrder": 5.0,
                    "ReorderPoint": 60.5,
                    "OrderUpTo": 110.5,
                    "SuggestedOrder": 86,
                }
            ],
        }
    )
    manifest = {
        "analysis_start": "2011-01-01",
        "analysis_end": "2011-12-31",
        "calendar_unknown_zero_days": 0,
        "provenance": {"input_name": "synthetic", "sha256": "synthetic"},
    }
    (folder / "workbook_data.json").write_text(json.dumps(data))
    (folder / "manifest.json").write_text(json.dumps(manifest))
    return data, export_excel(folder)


class PortabilityTests(unittest.TestCase):
    def test_controller_import_without_tk(self):
        code = """
import builtins
original = builtins.__import__
def blocked(name, *args, **kwargs):
    if name == 'tkinter' or name.startswith('tkinter.'):
        raise ImportError('Tk deliberately unavailable')
    return original(name, *args, **kwargs)
builtins.__import__ = blocked
from app import options
import test_pipeline
assert options('7','7','14','','{}')['lead_days'] == 7
"""
        run = subprocess.run(
            [sys.executable, "-c", code],
            cwd=Path(__file__).parent,
            capture_output=True,
            text=True,
        )
        self.assertEqual(run.returncode, 0, run.stderr)

    def test_verifier_uses_cli_folder_and_fails_missing_outputs(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / "custom"
            run = subprocess.run(
                [
                    sys.executable,
                    str(Path(__file__).with_name("verify_outputs.py")),
                    str(folder),
                ],
                capture_output=True,
                text=True,
            )
            self.assertNotEqual(run.returncode, 0)
            record = json.loads((folder / "verification.json").read_text())
            self.assertEqual(record["status"], "FAIL")
            self.assertEqual(record["output_folder"], str(folder.resolve()))

    def test_failed_verification_replaces_stale_pass(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "verification.json"
            path.write_text('{"status":"PASS"}')
            record = verify(tmp)
            self.assertEqual(json.loads(path.read_text())["status"], "FAIL")
            self.assertEqual(record["regression_tests"]["status"], "NOT_RUN")
            self.assertNotIn("regression_tests_passed", record)


class ReferenceTests(unittest.TestCase):
    def test_all_hw_candidates_against_independent_ets(self):
        rng = np.random.default_rng(123)
        examples = [
            np.full(70, 10.0),
            np.tile([10, 20, 30, 40, 50, 0, 60], 12),
            np.maximum(0, 20 + np.arange(121) * 0.2 + rng.normal(0, 3, 121)),
            rng.poisson(2, 180).astype(float),
        ]
        examples[-1][::31] += 150
        for series in examples:
            for params in GRIDS["Holt-Winters"]:
                for horizon in (1, 7, 14, 28):
                    with self.subTest(
                        length=len(series), params=params, horizon=horizon
                    ):
                        np.testing.assert_allclose(
                            hw(series, horizon, **params),
                            reference_hw(series, horizon, **params),
                            rtol=1e-10,
                            atol=1e-8,
                        )

    def test_hw_invalid_coefficients(self):
        with self.assertRaises(ValueError):
            hw(np.ones(40), 14, alpha=1.2)

    def test_calendar_imputation_cannot_see_future(self):
        s = pd.Series(
            np.arange(35, dtype=float), index=pd.date_range("2011-01-01", periods=35)
        )
        missing = [s.index[21]]
        changed = s.copy()
        changed.iloc[22:] = 1e9
        self.assertEqual(past_weekday_fill(s, missing).iloc[21], 7.0)
        self.assertEqual(
            past_weekday_fill(s, missing).iloc[21],
            past_weekday_fill(changed, missing).iloc[21],
        )


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.data, self.path = workbook_fixture(self.folder)

    def tearDown(self):
        self.temp.cleanup()

    def check(self, condition, label):
        self.assertTrue(condition, label)

    def test_complete_workbook_and_independent_formulas(self):
        result = check_workbook(self.path, self.data, self.check)
        self.assertEqual(result["status"], "PASS")
        self.assertEqual(result["native_engine_cache_status"], "NOT_AVAILABLE")
        self.assertEqual(result["formulas_evaluated"], 19)

    def test_missing_summary_is_a_failure(self):
        book = load_workbook(self.path)
        del book["Summary"]
        book.save(self.path)
        book.close()
        with self.assertRaises(AssertionError):
            check_workbook(self.path, self.data, self.check)

    def test_wrong_summary_value_is_a_failure(self):
        book = load_workbook(self.path)
        book["Summary"]["D5"] = "=999999"
        book.save(self.path)
        book.close()
        with self.assertRaises(AssertionError):
            check_workbook(self.path, self.data, self.check)

    def test_stock_and_on_order_changes_recalculate(self):
        book = load_workbook(self.path)
        sheet = book["Inventory"]
        columns = {c.value: c.column_letter for c in sheet[1]}
        for stock, on_order in [(0, 0), (20.2, 10.1), (200, 0), (0, 200)]:
            sheet[f"{columns['CurrentStock']}2"] = stock
            sheet[f"{columns['OnOrder']}2"] = on_order
            actual = formula_value(book, "Inventory", f"{columns['SuggestedOrder']}2")
            self.assertEqual(actual, max(0, math.ceil(110.5 - stock - on_order)))
        book.close()

    def test_missing_stock_is_not_silently_zero(self):
        book = load_workbook(self.path)
        sheet = book["Inventory"]
        columns = {c.value: c.column_letter for c in sheet[1]}
        sheet[f"{columns['CurrentStock']}2"] = None
        with self.assertRaises(ValueError):
            formula_value(book, "Inventory", f"{columns['SuggestedOrder']}2")
        book.close()

    def test_source_text_is_not_an_excel_formula(self):
        book = load_workbook(self.path)
        self.assertEqual(book["Products"]["B2"].data_type, "s")
        self.assertTrue(book["Products"]["B2"].value.startswith("'="))
        book.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
