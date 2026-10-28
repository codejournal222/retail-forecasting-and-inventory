import ast
import math
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from local_excel import REQUIRED_SHEETS, SHEET_ALIASES


def formula_value(book, sheet, address, visiting=None):
    visiting = set() if visiting is None else visiting
    key = (sheet, address)
    if key in visiting:
        raise ValueError("Circular workbook reference")
    value = book[sheet][address].value
    if not isinstance(value, str) or not value.startswith("="):
        if not isinstance(value, (int, float)) or not math.isfinite(value):
            raise ValueError(f"Non-numeric formula input: {sheet}!{address}")
        return float(value)
    visiting = visiting | {key}
    references = {}

    def reference(match):
        source = match.group(1) or sheet
        cell = match.group(2).replace("$", "")
        name = f"ref_{len(references)}"
        references[name] = formula_value(book, source, cell, visiting)
        return name

    expression = re.sub(r"(?:'([^']+)'!)?(\$?[A-Z]+\$?\d+)", reference, value[1:])

    def calculate(node):
        if isinstance(node, ast.Expression):
            return calculate(node.body)
        if isinstance(node, ast.Constant) and type(node.value) in (int, float):
            return node.value
        if isinstance(node, ast.Name) and node.id in references:
            return references[node.id]
        if isinstance(node, ast.BinOp):
            left, right = calculate(node.left), calculate(node.right)
            if isinstance(node.op, ast.Add):
                return left + right
            if isinstance(node.op, ast.Sub):
                return left - right
            if isinstance(node.op, ast.Mult):
                return left * right
            if isinstance(node.op, ast.Div):
                return left / right
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.USub):
            return -calculate(node.operand)
        if (
            isinstance(node, ast.Call)
            and isinstance(node.func, ast.Name)
            and not node.keywords
        ):
            args = [calculate(a) for a in node.args]
            if node.func.id == "MAX":
                return max(args)
            if node.func.id == "ROUNDUP" and len(args) == 2 and args[1] == 0:
                return math.copysign(math.ceil(abs(args[0])), args[0])
        raise ValueError("Unsupported workbook formula syntax")

    return calculate(ast.parse(expression, mode="eval"))


def check_workbook(path, data, check):
    workbook = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    formula_cells = []
    try:
        check(
            REQUIRED_SHEETS <= set(workbook.sheetnames),
            "Required workbook sheets are present",
        )
        for source, records in data.items():
            name = SHEET_ALIASES.get(source, source)
            check(name in workbook.sheetnames, f"Workbook sheet exists: {name}")
            sheet = workbook[name]
            if not records:
                check(sheet.max_row <= 1, f"Empty source stays empty: {name}")
                continue
            headers = list(records[0])
            check([c.value for c in sheet[1]] == headers, f"Workbook headers: {name}")
            check(sheet.max_row == len(records) + 1, f"Workbook row count: {name}")
            for row, record in enumerate(records, 2):
                for col, key in enumerate(headers, 1):
                    cell = sheet.cell(row, col)
                    expected = record[key]
                    if name == "Inventory" and key in {
                        "ReorderPoint",
                        "OrderUpTo",
                        "SuggestedOrder",
                    }:
                        continue
                    if key in {"Date", "Origin", "Start", "End"} and isinstance(
                        expected, str
                    ):
                        same = pd.Timestamp(cell.value) == pd.Timestamp(
                            expected
                        ).tz_localize(None)
                    elif isinstance(expected, (float, int)) and not isinstance(
                        expected, bool
                    ):
                        same = isinstance(cell.value, (float, int)) and np.isclose(
                            cell.value, expected, rtol=1e-9, atol=1e-8
                        )
                    else:
                        if isinstance(expected, str) and expected.startswith(
                            ("=", "+", "-", "@")
                        ):
                            expected = "'" + expected
                        same = cell.value == expected or (
                            expected == "" and cell.value is None
                        )
                    if not same:
                        check(
                            False, f"Workbook source mismatch: {name}!{cell.coordinate}"
                        )
            check(True, f"Workbook source values reconcile: {name}")
        inventory = workbook["Inventory"]
        headers = {c.value: get_column_letter(c.column) for c in inventory[1]}
        records = data["Replenishment"]
        for row, record in enumerate(records, 2):
            ref = lambda key: f"{headers[key]}{row}"
            expected_formulas = {
                "ReorderPoint": f"={ref('ForecastLeadUnits')}+{ref('SafetyStock')}",
                "OrderUpTo": f"={ref('ForecastProtectionUnits')}+{ref('SafetyStock')}",
                "SuggestedOrder": f"=MAX(0,ROUNDUP({ref('OrderUpTo')}-{ref('CurrentStock')}-{ref('OnOrder')},0))",
            }
            for key, expected in expected_formulas.items():
                check(
                    inventory[ref(key)].value == expected,
                    f"Inventory formula references: {ref(key)}",
                )
                expected_value = {
                    "ReorderPoint": record["ForecastLeadUnits"] + record["SafetyStock"],
                    "OrderUpTo": record["ForecastProtectionUnits"]
                    + record["SafetyStock"],
                    "SuggestedOrder": max(
                        0,
                        math.ceil(
                            record["ForecastProtectionUnits"]
                            + record["SafetyStock"]
                            - record["CurrentStock"]
                            - record["OnOrder"]
                        ),
                    ),
                }[key]
                check(
                    np.isclose(
                        formula_value(workbook, "Inventory", ref(key)), expected_value
                    ),
                    f"Independent inventory formula value: {ref(key)}",
                )
        for row, record in enumerate(data["Model_Comparison"], 5):
            check(
                workbook["Summary"].cell(row, 1).value == record["Model"],
                f"Summary model: row {row}",
            )
            for col, source, metric in [
                (2, "Model_Comparison", "MAE"),
                (3, "Protection_Comparison", "MAE"),
                (4, "Inventory_Comparison", "TotalCost"),
                (5, "Inventory_Comparison", "FillRate"),
            ]:
                expected = next(
                    r[metric] for r in data[source] if r["Model"] == record["Model"]
                )
                address = f"{get_column_letter(col)}{row}"
                check(
                    workbook["Summary"][address].data_type == "f",
                    f"Summary links to source: {address}",
                )
                check(
                    np.isclose(formula_value(workbook, "Summary", address), expected),
                    f"Summary value reconciles: {address}",
                )
        missing_cache = 0
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    check_value = cached[sheet.title][cell.coordinate]
                    if cell.data_type == "e" or check_value.data_type == "e":
                        check(False, f"Excel error: {sheet.title}!{cell.coordinate}")
                    if cell.data_type == "f":
                        expected = formula_value(workbook, sheet.title, cell.coordinate)
                        formula_cells.append((sheet.title, cell.coordinate))
                        if check_value.value is None:
                            missing_cache += 1
                        else:
                            check(
                                np.isclose(check_value.value, expected),
                                f"Cached formula value: {sheet.title}!{cell.coordinate}",
                            )
        check(
            True,
            "Every workbook formula evaluated with restricted independent evaluator",
        )
        return {
            "status": "PASS",
            "formulas_evaluated": len(formula_cells),
            "native_engine_cache_status": (
                "NOT_AVAILABLE" if missing_cache else "VALUES_VERIFIED"
            ),
            "missing_formula_caches": missing_cache,
            "engine_limit": "Restricted evaluator, not Microsoft Excel. Open in Excel to recalculate native caches.",
        }
    finally:
        workbook.close()
        cached.close()
